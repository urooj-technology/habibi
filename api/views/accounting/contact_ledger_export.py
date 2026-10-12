"""
Excel / PDF builders for customer and supplier ledger exports.
Payload shape matches CustomerLedgerAPIView / SupplierLedgerAPIView JSON responses.
"""

import io
import csv
import re
from datetime import datetime
from decimal import Decimal

from django.http import HttpResponse


def _d2(val) -> float:
    try:
        return float(Decimal(str(val or "0")).quantize(Decimal("0.01")))
    except Exception:
        return 0.0


def _safe_filename_fragment(s: str) -> str:
    s = (s or "").strip()
    return re.sub(r"[^\w\-.]+", "_", s)[:80] if s else ""


def _paragraph_escape(s: str) -> str:
    return (
        str(s or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def _export_sections(data: dict) -> list[dict]:
    """
    Prefer per-currency sections when All currencies is selected.
    Fall back to a single flat section for one-currency responses.
    """
    sections = data.get("sections") or []
    if sections:
        return sections
    display_ccy = data.get("currency") or ""
    if display_ccy == "MIXED":
        display_ccy = ""
    return [
        {
            "currency": display_ccy or data.get("currency_filter") or "",
            "opening_balance": data.get("opening_balance", "0"),
            "closing_balance": data.get("closing_balance", "0"),
            "period_debit": data.get("period_debit", "0"),
            "period_credit": data.get("period_credit", "0"),
            "lines": data.get("lines") or [],
        }
    ]


def excel_contact_ledger_response(
    data: dict, *, ledger_label: str, contact_name: str, entity_key: str, entity_id: int
) -> HttpResponse:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except Exception as e:
        return HttpResponse(
            f"Excel export dependency missing: {e}",
            status=500,
            content_type="text/plain",
        )

    date_from = data.get("from", "")
    date_to = data.get("to", "")
    sections = _export_sections(data)
    multi = len(sections) > 1

    wb = Workbook()
    # Remove default sheet; create one sheet per currency section.
    default_ws = wb.active
    wb.remove(default_ws)

    title_font = Font(bold=True, size=16, color="1F2937")
    subtitle_font = Font(size=10, color="6B7280")
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="111827")
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    generated_at = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")
    subtitle_contact = contact_name if contact_name else f"#{entity_id}"

    for section in sections:
        ccy = section.get("currency") or "AFN"
        lines = section.get("lines") or []
        opening = section.get("opening_balance", "0")
        closing = section.get("closing_balance", "0")
        sheet_title = ccy if multi else "Ledger"
        ws = wb.create_sheet(title=sheet_title[:31])

        ws.merge_cells("A1:J1")
        ws["A1"] = f"{ledger_label} — {subtitle_contact}"
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("A2:J2")
        ws["A2"] = (
            f"Period: {date_from} — {date_to}  •  Currency: {ccy}"
            f"  •  Opening: {_d2(opening):,.2f} {ccy}"
            f"  •  Closing: {_d2(closing):,.2f} {ccy}"
        )
        ws["A2"].font = subtitle_font
        ws["A2"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("A3:J3")
        ws["A3"] = f"Generated at: {generated_at}"
        ws["A3"].font = subtitle_font
        ws["A3"].alignment = Alignment(horizontal="left", vertical="center")

        ws.append([])

        headers = [
            "Date",
            "Voucher No",
            "Reference No",
            "Transaction Type",
            "Description",
            f"Debit ({ccy})",
            f"Credit ({ccy})",
            f"Balance ({ccy})",
            "Created By",
            "Status",
        ]
        ws.append(headers)
        header_row = ws.max_row
        for col_idx, name in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col_idx, value=name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = border

        data_start = ws.max_row + 1
        for ln in lines:
            ws.append(
                [
                    ln.get("date") or ln.get("entry_date", ""),
                    ln.get("voucher_no") or ln.get("reference") or "",
                    ln.get("reference_no") or ln.get("reference") or "",
                    ln.get("transaction_type") or ln.get("source_type") or "",
                    ln.get("description") or ln.get("line_memo") or "",
                    _d2(ln.get("debit") or ln.get("debit_home")),
                    _d2(ln.get("credit") or ln.get("credit_home")),
                    _d2(ln.get("running_balance")),
                    ln.get("created_by") or "",
                    ln.get("status") or "",
                ]
            )

        data_end = ws.max_row if lines else header_row

        money_cols = {6, 7, 8}
        for row in range(data_start, data_end + 1):
            for col in range(1, 11):
                c = ws.cell(row=row, column=col)
                c.border = border
                if col == 1 or col == 3:
                    c.alignment = align_center
                elif col in money_cols:
                    c.alignment = align_right
                    c.number_format = "#,##0.00"
                else:
                    c.alignment = align_left

        ws.freeze_panes = f"A{data_start}"
        if lines:
            ws.auto_filter.ref = f"A{header_row}:J{data_end}"

        widths = {1: 12, 2: 18, 3: 18, 4: 22, 5: 36, 6: 16, 7: 16, 8: 18, 9: 18, 10: 14}
        for col_idx, w in widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    name_part = _safe_filename_fragment(contact_name)
    ts = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    fname = (
        f"{entity_key}-{entity_id}-{date_from}-{date_to}-{ts}.xlsx"
        if not name_part
        else f"{entity_key}-{name_part}-{date_from}-{date_to}-{ts}.xlsx"
    )

    resp = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f'attachment; filename="{fname}"'
    return resp


def pdf_contact_ledger_response(
    data: dict, *, ledger_label: str, contact_name: str, entity_key: str, entity_id: int
) -> HttpResponse:
    try:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import (
            SimpleDocTemplate,
            Table,
            TableStyle,
            Paragraph,
            Spacer,
            PageBreak,
        )
    except Exception as e:
        return HttpResponse(
            f"PDF export dependency missing: {e}",
            status=500,
            content_type="text/plain",
        )

    date_from = data.get("from", "")
    date_to = data.get("to", "")
    sections = _export_sections(data)

    subtitle_contact = contact_name if contact_name else f"#{entity_id}"
    generated_at = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=18,
        rightMargin=18,
        topMargin=18,
        bottomMargin=22,
        title=ledger_label,
    )
    styles = getSampleStyleSheet()

    page_w, _page_h = landscape(A4)
    usable_w = page_w - doc.leftMargin - doc.rightMargin
    col_widths = [
        usable_w * 0.085,
        usable_w * 0.1,
        usable_w * 0.095,
        usable_w * 0.12,
        usable_w * 0.22,
        usable_w * 0.085,
        usable_w * 0.085,
        usable_w * 0.09,
        usable_w * 0.065,
        usable_w * 0.055,
    ]

    story = []
    for idx, section in enumerate(sections):
        if idx:
            story.append(PageBreak())

        ccy = section.get("currency") or "AFN"
        lines = section.get("lines") or []
        opening = section.get("opening_balance", "0")
        closing = section.get("closing_balance", "0")

        story.append(
            Paragraph(
                _paragraph_escape(f"{ledger_label}: {subtitle_contact} ({ccy})"),
                styles["Title"],
            )
        )
        story.append(
            Paragraph(
                _paragraph_escape(
                    f"Period {date_from} — {date_to} | Opening: {opening} {ccy} | Closing: {closing} {ccy}"
                ),
                styles["Normal"],
            )
        )
        story.append(Paragraph(_paragraph_escape(f"Generated at: {generated_at}"), styles["Normal"]))
        story.append(Spacer(1, 8))

        tbl_header = [
            "Date",
            "Voucher",
            "Ref.",
            "Type",
            "Description",
            f"Debit ({ccy})",
            f"Credit ({ccy})",
            f"Balance ({ccy})",
            "Created By",
            "Status",
        ]
        table_rows = [tbl_header]
        for ln in lines:
            table_rows.append(
                [
                    str(ln.get("date") or ln.get("entry_date") or ""),
                    str(ln.get("voucher_no") or "-")[:24],
                    str(ln.get("reference_no") or "-")[:24],
                    str(ln.get("transaction_type") or "-")[:24],
                    str(ln.get("description") or "-")[:52],
                    str(ln.get("debit") or ln.get("debit_home") or "0.00"),
                    str(ln.get("credit") or ln.get("credit_home") or "0.00"),
                    str(ln.get("running_balance") or "0.00"),
                    str(ln.get("created_by") or "-")[:18],
                    str(ln.get("status") or "-")[:16],
                ]
            )

        table = Table(table_rows, repeatRows=1, colWidths=col_widths)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#111827")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 8),
                    ("FONTSIZE", (0, 1), (-1, -1), 7),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D1D5DB")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ("ALIGN", (0, 1), (0, -1), "CENTER"),
                    ("ALIGN", (1, 1), (3, -1), "CENTER"),
                    ("ALIGN", (5, 1), (7, -1), "RIGHT"),
                ]
            )
        )
        story.append(table)

    doc.build(story)

    pdf = buf.getvalue()
    buf.close()

    name_part = _safe_filename_fragment(contact_name)
    ts = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    fname = (
        f"{entity_key}-{entity_id}-{date_from}-{date_to}-{ts}.pdf"
        if not name_part
        else f"{entity_key}-{name_part}-{date_from}-{date_to}-{ts}.pdf"
    )

    resp = HttpResponse(pdf, content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="{fname}"'
    return resp


def csv_contact_ledger_response(
    data: dict, *, ledger_label: str, contact_name: str, entity_key: str, entity_id: int
) -> HttpResponse:
    date_from = data.get("from", "")
    date_to = data.get("to", "")
    sections = _export_sections(data)
    name_part = _safe_filename_fragment(contact_name)
    ts = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    fname = (
        f"{entity_key}-{entity_id}-{date_from}-{date_to}-{ts}.csv"
        if not name_part
        else f"{entity_key}-{name_part}-{date_from}-{date_to}-{ts}.csv"
    )

    resp = HttpResponse(content_type="text/csv")
    resp["Content-Disposition"] = f'attachment; filename="{fname}"'
    writer = csv.writer(resp)
    writer.writerow([ledger_label, contact_name or f"#{entity_id}"])
    writer.writerow(["Period", date_from, date_to])

    for section in sections:
        ccy = section.get("currency") or ""
        writer.writerow([])
        writer.writerow(
            [
                f"Currency: {ccy}",
                f"Opening: {section.get('opening_balance', '0')}",
                f"Closing: {section.get('closing_balance', '0')}",
            ]
        )
        writer.writerow(
            [
                "Date",
                "Voucher No",
                "Reference No",
                "Transaction Type",
                "Description",
                f"Debit ({ccy})",
                f"Credit ({ccy})",
                f"Running Balance ({ccy})",
                "Created By",
                "Status",
                "Currency",
            ]
        )
        for ln in section.get("lines") or []:
            writer.writerow(
                [
                    ln.get("date") or ln.get("entry_date") or "",
                    ln.get("voucher_no") or "",
                    ln.get("reference_no") or "",
                    ln.get("transaction_type") or "",
                    ln.get("description") or ln.get("line_memo") or "",
                    ln.get("debit") or ln.get("debit_home") or "0.00",
                    ln.get("credit") or ln.get("credit_home") or "0.00",
                    ln.get("running_balance") or "0.00",
                    ln.get("created_by") or "",
                    ln.get("status") or "",
                    ln.get("currency") or ccy,
                ]
            )
    return resp
