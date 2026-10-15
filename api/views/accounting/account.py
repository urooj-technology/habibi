from decimal import Decimal

import io
from datetime import datetime

from django.http import HttpResponse
from django_filters.rest_framework import DjangoFilterBackend
from django.db.models import DecimalField, Sum, Value
from django.db.models.functions import Coalesce

from api.services.accounting.balances import annotate_account_balance_totals
from rest_framework.filters import SearchFilter
from rest_framework.permissions import IsAuthenticated
from account.permissions import ModuleCRUDPermission
from rest_framework import viewsets
from rest_framework.decorators import action

from api.common.pdf_export import HEADER_BG, apply_unicode_table_fonts, pdf_header_cell, pdf_paragraph
from api.common.pdf_fonts import ensure_pdf_fonts
from api.models.accounting.account import AccountingAccount
from api.serializers.accounting.account import AccountingAccountSerializer
from api.views.data.pagination import StandardResultsPagination


class AccountingAccountViewSet(viewsets.ModelViewSet):
    """Chart of accounts (AFN home currency on journal lines)."""

    queryset = AccountingAccount.objects.all().order_by("code")
    serializer_class = AccountingAccountSerializer
    permission_classes = [ModuleCRUDPermission("accounting")]
    pagination_class = StandardResultsPagination
    filter_backends = [DjangoFilterBackend, SearchFilter]
    filterset_fields = ["account_type", "role", "currency"]
    search_fields = ["code", "name"]

    def get_queryset(self):
        qs = super().get_queryset()
        out_f = DecimalField(max_digits=18, decimal_places=4)
        zero_f = Value(Decimal("0.0000"), output_field=out_f)
        qs = annotate_account_balance_totals(qs)
        debit_f = Coalesce(
            Sum("journal_lines__debit_foreign", output_field=out_f),
            zero_f,
            output_field=out_f,
        )
        credit_f = Coalesce(
            Sum("journal_lines__credit_foreign", output_field=out_f),
            zero_f,
            output_field=out_f,
        )
        return qs.annotate(
            _debit_foreign_raw=debit_f,
            _credit_foreign_raw=credit_f,
        )

    def _export_queryset(self):
        qs = self.filter_queryset(self.get_queryset())
        return qs.order_by("code")

    @action(detail=False, methods=["get"], url_path="export/excel")
    def export_excel(self, request, *args, **kwargs):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
            from openpyxl.utils import get_column_letter
        except Exception as e:  # pragma: no cover
            return HttpResponse(
                f"Excel export dependency missing: {e}",
                status=500,
                content_type="text/plain",
            )

        rows = AccountingAccountSerializer(self._export_queryset(), many=True).data

        wb = Workbook()
        ws = wb.active
        ws.title = "Chart of Accounts"

        generated_at = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")

        # Styles
        title_font = Font(bold=True, size=16, color="1F2937")
        subtitle_font = Font(size=10, color="6B7280")
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="111827")
        thin = Side(style="thin", color="D1D5DB")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        align_center = Alignment(horizontal="center", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")

        # Title rows
        ws.merge_cells("A1:H1")
        ws["A1"] = "Chart of Accounts"
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("A2:H2")
        ws["A2"] = f"Generated at: {generated_at}"
        ws["A2"].font = subtitle_font
        ws["A2"].alignment = Alignment(horizontal="left", vertical="center")

        ws.append([])  # row 3 spacer

        headers = [
            "Code",
            "Account Name",
            "Type",
            "Role",
            "Active",
            "Debit",
            "Credit",
            "Balance",
        ]
        ws.append(headers)  # row 4

        header_row = ws.max_row
        for col_idx, name in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col_idx, value=name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = border

        def d2(v):
            try:
                return float(Decimal(str(v or "0")).quantize(Decimal("0.01")))
            except Exception:
                return 0.0

        data_start = ws.max_row + 1
        for r in rows:
            ws.append(
                [
                    r.get("code", ""),
                    r.get("name", ""),
                    r.get("account_type", ""),
                    r.get("role", ""),
                    "Yes" if r.get("is_active") else "No",
                    d2(r.get("total_debit")),
                    d2(r.get("total_credit")),
                    d2(r.get("balance")),
                ]
            )

        data_end = ws.max_row

        # Formats & borders for data
        money_cols = [6, 7, 8]
        for row in range(data_start, data_end + 1):
            for col in range(1, 9):
                c = ws.cell(row=row, column=col)
                c.border = border
                if col in (1, 5):
                    c.alignment = align_center
                elif col in money_cols:
                    c.alignment = align_right
                    c.number_format = "#,##0.00"
                else:
                    c.alignment = align_left

        # Freeze panes under header row
        ws.freeze_panes = f"A{data_start}"

        # Auto-filter on header
        ws.auto_filter.ref = f"A{header_row}:H{data_end}"

        # Column widths (balanced for readability)
        widths = {
            1: 10,   # code
            2: 34,   # name
            3: 14,   # type
            4: 16,   # role
            5: 10,   # active
            6: 16,
            7: 16,
            8: 16,
        }
        for col_idx, w in widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = w

        out = io.BytesIO()
        wb.save(out)
        out.seek(0)

        fname = f"chart-of-accounts_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
        resp = HttpResponse(
            out.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        resp["Content-Disposition"] = f'attachment; filename="{fname}"'
        return resp

    @action(detail=False, methods=["get"], url_path="export/pdf")
    def export_pdf(self, request, *args, **kwargs):
        try:
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib import colors
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Spacer
        except Exception as e:  # pragma: no cover
            return HttpResponse(
                f"PDF export dependency missing: {e}",
                status=500,
                content_type="text/plain",
            )

        rows = AccountingAccountSerializer(self._export_queryset(), many=True).data
        font_regular, font_bold = ensure_pdf_fonts()
        styles = getSampleStyleSheet()

        buf = io.BytesIO()
        doc = SimpleDocTemplate(
            buf,
            pagesize=landscape(A4),
            leftMargin=20,
            rightMargin=20,
            topMargin=20,
            bottomMargin=20,
            title="Chart of Accounts",
        )

        generated_at = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")
        title = pdf_paragraph("Chart of Accounts", styles["Title"], bold=True)
        subtitle = pdf_paragraph(f"Generated at: {generated_at}")

        headers = ["Code", "Name", "Type", "Role", "Active", "Debit", "Credit", "Balance"]
        data = [[pdf_header_cell(h) for h in headers]]
        for r in rows:
            data.append(
                [
                    pdf_paragraph(r.get("code", "")),
                    pdf_paragraph(r.get("name", "")),
                    pdf_paragraph(r.get("account_type", "")),
                    pdf_paragraph(r.get("role", "")),
                    pdf_paragraph("Yes" if r.get("is_active") else "No"),
                    pdf_paragraph(r.get("total_debit", "0.00")),
                    pdf_paragraph(r.get("total_credit", "0.00")),
                    pdf_paragraph(r.get("balance", "0.00")),
                ]
            )

        page_w, _page_h = landscape(A4)
        usable_w = page_w - doc.leftMargin - doc.rightMargin
        col_widths = [
            usable_w * 0.08,
            usable_w * 0.34,
            usable_w * 0.12,
            usable_w * 0.14,
            usable_w * 0.08,
            usable_w * 0.08,
            usable_w * 0.08,
            usable_w * 0.08,
        ]

        table = Table(data, repeatRows=1, colWidths=col_widths)
        table.setStyle(
            TableStyle(
                apply_unicode_table_fonts(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), HEADER_BG),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, 0), font_bold),
                        ("FONTNAME", (0, 1), (-1, -1), font_regular),
                        ("FONTSIZE", (0, 0), (-1, 0), 9),
                        ("FONTSIZE", (0, 1), (-1, -1), 8),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D1D5DB")),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        (
                            "ROWBACKGROUNDS",
                            (0, 1),
                            (-1, -1),
                            [colors.white, colors.HexColor("#F9FAFB")],
                        ),
                        ("LEFTPADDING", (0, 0), (-1, -1), 6),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                        ("ALIGN", (0, 0), (0, -1), "CENTER"),
                        ("ALIGN", (4, 0), (4, -1), "CENTER"),
                        ("ALIGN", (5, 0), (-1, -1), "RIGHT"),
                    ]
                )
            )
        )

        doc.build([title, subtitle, Spacer(1, 10), table])
        pdf = buf.getvalue()
        buf.close()

        fname = f"chart-of-accounts_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.pdf"
        resp = HttpResponse(pdf, content_type="application/pdf")
        resp["Content-Disposition"] = f'attachment; filename="{fname}"'
        return resp
