"""Excel / PDF builders for stock inventory exports."""

import io
from datetime import datetime
from decimal import Decimal

from django.http import HttpResponse

from api.common.pdf_export import HEADER_BG, apply_unicode_table_fonts, pdf_header_cell, pdf_paragraph
from api.common.pdf_fonts import ensure_pdf_fonts


def _dec(value, places="0.01"):
    try:
        return float(Decimal(str(value or "0")).quantize(Decimal(places)))
    except Exception:
        return 0.0


def build_stock_export_rows(queryset):
    rows = []
    for inv in queryset.select_related("product", "product__category"):
        product = inv.product
        rows.append(
            {
                "id": inv.id,
                "product": product.name if product else "",
                "barcode": (product.barcode if product else "") or "",
                "category": (
                    product.category.name
                    if product and getattr(product, "category_id", None)
                    else ""
                ),
                "quantity": _dec(inv.quantity),
                "buy_price": _dec(inv.buy_price),
                "sell_price": _dec(inv.sell_price),
                "min_quantity": _dec(inv.min_quantity),
            }
        )
    return rows


def excel_stock_inventory_response(rows) -> HttpResponse:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except Exception as e:
        return HttpResponse(
            f"Excel export dependency missing: {e}",
            status=500,
            content_type="text/plain",
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "Stock"

    generated_at = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")
    title_font = Font(bold=True, size=16, color="1F2937")
    subtitle_font = Font(size=10, color="6B7280")
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="111827")
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    headers = [
        "ID",
        "Product",
        "Barcode",
        "Category",
        "Quantity",
        "Buy Price",
        "Sell Price",
        "Min Quantity",
    ]

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws["A1"] = "Stock Inventory"
    ws["A1"].font = title_font
    ws["A1"].alignment = align_left

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws["A2"] = f"Generated: {generated_at}"
    ws["A2"].font = subtitle_font

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = border

    for r_idx, row in enumerate(rows, start=5):
        values = [
            row["id"],
            row["product"],
            row["barcode"],
            row["category"],
            row["quantity"],
            row["buy_price"],
            row["sell_price"],
            row["min_quantity"],
        ]
        for c_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.border = border
            if c_idx >= 5:
                cell.alignment = align_right
                if c_idx >= 5:
                    cell.number_format = "#,##0.00"
            else:
                cell.alignment = align_left

    widths = {1: 8, 2: 28, 3: 16, 4: 16, 5: 12, 6: 12, 7: 12, 8: 12}
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="stock_inventory.xlsx"'
    return response


def pdf_stock_inventory_response(rows) -> HttpResponse:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Spacer
    except Exception as e:
        return HttpResponse(
            f"PDF export dependency missing: {e}",
            status=500,
            content_type="text/plain",
        )

    font_regular, font_bold = ensure_pdf_fonts()
    styles = getSampleStyleSheet()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=18,
        rightMargin=18,
        topMargin=18,
        bottomMargin=18,
        title="Stock Inventory",
    )
    story = [
        pdf_paragraph("Stock Inventory", styles["Title"], bold=True),
        Spacer(1, 12),
    ]

    headers = [
        "ID",
        "Product",
        "Barcode",
        "Category",
        "Qty",
        "Buy",
        "Sell",
        "Min",
    ]
    data = [[pdf_header_cell(h) for h in headers]]
    for row in rows:
        data.append(
            [
                pdf_paragraph(row["id"]),
                pdf_paragraph(row["product"]),
                pdf_paragraph(row["barcode"]),
                pdf_paragraph(row["category"]),
                pdf_paragraph(f"{row['quantity']:.2f}"),
                pdf_paragraph(f"{row['buy_price']:.2f}"),
                pdf_paragraph(f"{row['sell_price']:.2f}"),
                pdf_paragraph(f"{row['min_quantity']:.2f}"),
            ]
        )

    page_w, _ = landscape(A4)
    usable = page_w - 36
    col_widths = [usable * w for w in (0.06, 0.24, 0.14, 0.14, 0.1, 0.11, 0.11, 0.1)]
    table = Table(data, repeatRows=1, colWidths=col_widths)
    table.setStyle(
        TableStyle(
            apply_unicode_table_fonts(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), HEADER_BG),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("FONTNAME", (0, 0), (-1, 0), font_bold),
                    ("FONTNAME", (0, 1), (-1, -1), font_regular),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [colors.white, colors.HexColor("#F3F4F6")],
                    ),
                    ("ALIGN", (4, 1), (-1, -1), "RIGHT"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
    )
    story.append(table)
    doc.build(story)
    buf.seek(0)
    response = HttpResponse(buf.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="stock_inventory.pdf"'
    return response
